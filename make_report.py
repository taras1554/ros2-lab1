import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SEARCH_DIRS = [
    ".",
    os.path.expanduser("~/Зображення/Знімки екрана"),
    os.path.expanduser("~/Зображення"),
    os.path.expanduser("~/Pictures/Screenshots"),
    os.path.expanduser("~/Pictures"),
    os.path.expanduser("~/Завантаження"),
    os.path.expanduser("~/Downloads"),
]

def find_image(filename):
    for d in SEARCH_DIRS:
        path = os.path.join(d, filename)
        if os.path.exists(path):
            return path
    return None

def set_run_font(run, name="Times New Roman", size=14, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

def add_p(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, size=14, space_after=6, indent=0.5):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if indent:
        p.paragraph_format.first_line_indent = Inches(indent)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p

def add_h(doc, text, size=14):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True)
    return p

def add_code(doc, title, code_text):
    add_p(doc, title, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=12, space_after=2, indent=0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(code_text.strip())
    set_run_font(run, name="Courier New", size=10)

def add_img(doc, filename, caption):
    path = find_image(filename)
    if path:
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(path, width=Inches(6.2))
        print(f"[OK] Додано скріншот: {path}")
    else:
        add_p(doc, f"[Місце для зображення: {filename}]", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=12, indent=0)
        print(f"[!] Не знайдено файл: {filename}")
    add_p(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=12, space_after=10, indent=0)

doc = Document()
for s in doc.sections:
    s.top_margin = Inches(0.79)
    s.bottom_margin = Inches(0.79)
    s.left_margin = Inches(0.98)
    s.right_margin = Inches(0.59)

# Титульна сторінка
add_p(doc, "НАЦІОНАЛЬНИЙ ТЕХНІЧНИЙ УНІВЕРСИТЕТ УКРАЇНИ\n«КИЇВСЬКИЙ ПОЛІТЕХНІЧНИЙ ІНСТИТУТ імені ІГОРЯ СІКОРСЬКОГО»", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, indent=0)
add_p(doc, "Факультет прикладної математики (ФПСПМ)", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=72, indent=0)
add_p(doc, "ЛАБОРАТОРНА РОБОТА №1", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, space_after=6, indent=0)
add_p(doc, "Тема: «Вступ до ROS 2, налаштування середовища та розробка вузлів (Nodes)»", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, space_after=96, indent=0)
add_p(doc, "Виконав:\nстудент групи КВ-34\nТурчик Тарас", align=WD_ALIGN_PARAGRAPH.RIGHT, size=14, space_after=120, indent=0)
add_p(doc, "Київ – 2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, indent=0)
doc.add_page_break()

# Розділ 1
add_h(doc, "1. Тема та мета роботи")
add_p(doc, "Тема: Вступ до ROS. Вузли.")
add_p(doc, "Мета роботи: Опанувати засоби контейнеризації середовища розробки (Docker), навчитися створювати та збирати ROS 2 пакети за допомогою утиліти colcon, проектувати вузли у вигляді компонентів (Components) для роботи у межах єдиного процесу-виконавця (Executor) та реалізовувати вузли з керованим життєвим циклом (Lifecycle Nodes).")

# Розділ 2
add_h(doc, "2. Виконання практичних завдань")
add_h(doc, "Завдання 1.1: Налаштування середовища та робочого простору")
add_p(doc, "Для ізоляції середовища розробки підготовлено Dockerfile на основі офіційного образу osrf/ros:humble-desktop із встановленням python3-colcon-common-extensions, ros-humble-rclcpp-components, ros-humble-rclcpp-lifecycle та ros-humble-lifecycle-msgs. Для автоматичного сорсингу системного та локального setup.bash створено скрипт entrypoint.sh.")

add_code(doc, "Лістинг 1.1 – Файл Dockerfile:", """FROM osrf/ros:humble-desktop

RUN apt-get update && apt-get install -y \\
    python3-colcon-common-extensions \\
    ros-humble-rclcpp-components \\
    ros-humble-rclcpp-lifecycle \\
    ros-humble-lifecycle-msgs \\
    && rm -rf /var/lib/apt/lists/*

WORKDIR /ros2_ws
RUN mkdir -p /ros2_ws/src

COPY entrypoint.sh /entrypoint.sh
RUN chmod +x /entrypoint.sh

ENTRYPOINT ["/entrypoint.sh"]
CMD ["bash"]""")

add_code(doc, "Лістинг 1.2 – Файл entrypoint.sh:", """#!/bin/bash
set -e

source /opt/ros/$ROS_DISTRO/setup.bash

if [ -f /ros2_ws/install/setup.bash ]; then
    source /ros2_ws/install/setup.bash
fi

exec "$@\"""" )

add_code(doc, "Лістинг 1.3 – Файл CMakeLists.txt пакета lab1_pkg:", """cmake_minimum_required(VERSION 3.8)
project(lab1_pkg)

if(CMAKE_COMPILER_IS_GNUCXX OR CMAKE_CXX_COMPILER_ID MATCHES "Clang")
  add_compile_options(-Wall -Wextra -Wpedantic)
endif()

find_package(ament_cmake REQUIRED)
find_package(rclcpp REQUIRED)
find_package(rclcpp_components REQUIRED)
find_package(rclcpp_lifecycle REQUIRED)
find_package(lifecycle_msgs REQUIRED)

# Завдання 1.2: Спільна бібліотека компонента
add_library(my_component SHARED src/my_component.cpp)
ament_target_dependencies(my_component rclcpp rclcpp_components)
rclcpp_components_register_nodes(my_component "lab1_ns::MyComponent")

# Завдання 1.3: Керовані вузли та супервізор
add_executable(managed_nodes_exec src/managed_node.cpp)
ament_target_dependencies(managed_nodes_exec rclcpp rclcpp_lifecycle lifecycle_msgs)

add_executable(lifecycle_supervisor src/lifecycle_supervisor.cpp)
ament_target_dependencies(lifecycle_supervisor rclcpp lifecycle_msgs)

install(TARGETS my_component
  ARCHIVE DESTINATION lib
  LIBRARY DESTINATION lib
  RUNTIME DESTINATION bin
)

install(TARGETS managed_nodes_exec lifecycle_supervisor
  DESTINATION lib/${PROJECT_NAME}
)

ament_package()""")

add_h(doc, "Завдання 1.2: Розробка вузла-компонента (Component Node)")
add_p(doc, "У файлі src/my_component.cpp створено клас MyComponent у просторі імен lab1_ns, успадкований від rclcpp::Node, який приймає конструктор rclcpp::NodeOptions. Вузол зареєстровано як компонент макросом RCLCPP_COMPONENTS_REGISTER_NODE(lab1_ns::MyComponent).")

add_code(doc, "Лістинг 1.4 – Вихідний код src/my_component.cpp:", """#include "rclcpp/rclcpp.hpp"
#include "rclcpp_components/register_node_macro.hpp"

namespace lab1_ns
{

class MyComponent : public rclcpp::Node
{
public:
  explicit MyComponent(const rclcpp::NodeOptions & options)
  : Node("my_component_node", options)
  {
    RCLCPP_INFO(this->get_logger(), "Вузол-компонент MyComponent успішно ініціалізовано!");
  }
};

}  

RCLCPP_COMPONENTS_REGISTER_NODE(lab1_ns::MyComponent)""")

add_p(doc, "Виконано збирання пакета командою colcon build --symlink-install (час збирання 9.28 с), сорсинг install/setup.bash та запуск процесу component_container. У другому терміналі через ros2 component load /ComponentManager lab1_pkg lab1_ns::MyComponent динамічно завантажено бібліотеку libmy_component.so.")

add_img(doc, "Знімок екрана з 2026-10-15 16-00-03.png", "Рисунок 1.1 – Збирання пакета lab1_pkg та ініціалізація компонента у component_container")
add_img(doc, "Знімок екрана з 2026-10-15 16-00-17.png", "Рисунок 1.2 – Завантаження компонента lab1_ns::MyComponent у другому терміналі")

add_h(doc, "Завдання 1.3: Розробка керованих вузлів (Lifecycle Nodes) та супервізора")
add_p(doc, "У файлі src/managed_node.cpp реалізовано клас ManagedNode на базі rclcpp_lifecycle::LifecycleNode та запущено два екземпляри (managed_node_1, managed_node_2) у SingleThreadedExecutor. У файлі src/lifecycle_supervisor.cpp реалізовано супервізор, який через сервіс change_state виконує переходи TRANSITION_CONFIGURE та TRANSITION_ACTIVATE.")

add_code(doc, "Лістинг 1.5 – Вихідний код src/managed_node.cpp:", """#include <memory>
#include "rclcpp/rclcpp.hpp"
#include "rclcpp_lifecycle/lifecycle_node.hpp"

using CallbackReturn = rclcpp_lifecycle::node_interfaces::LifecycleNodeInterface::CallbackReturn;

class ManagedNode : public rclcpp_lifecycle::LifecycleNode
{
public:
  explicit ManagedNode(const std::string & node_name)
  : rclcpp_lifecycle::LifecycleNode(node_name)
  {
    RCLCPP_INFO(get_logger(), "[%s] створено (стан: Unconfigured).", node_name.c_str());
  }

  CallbackReturn on_configure(const rclcpp_lifecycle::State &) override
  {
    RCLCPP_INFO(get_logger(), "[%s] on_configure(): Конфігурація...", get_name());
    return CallbackReturn::SUCCESS;
  }

  CallbackReturn on_activate(const rclcpp_lifecycle::State & state) override
  {
    LifecycleNode::on_activate(state);
    RCLCPP_INFO(get_logger(), "[%s] on_activate(): Вузол АКТИВОВАНО!", get_name());
    return CallbackReturn::SUCCESS;
  }

  CallbackReturn on_deactivate(const rclcpp_lifecycle::State & state) override
  {
    LifecycleNode::on_deactivate(state);
    RCLCPP_INFO(get_logger(), "[%s] on_deactivate(): Деактивовано.", get_name());
    return CallbackReturn::SUCCESS;
  }

  CallbackReturn on_cleanup(const rclcpp_lifecycle::State &) override
  {
    RCLCPP_INFO(get_logger(), "[%s] on_cleanup(): Очищення ресурсів.", get_name());
    return CallbackReturn::SUCCESS;
  }

  CallbackReturn on_shutdown(const rclcpp_lifecycle::State &) override
  {
    RCLCPP_INFO(get_logger(), "[%s] on_shutdown(): Завершення роботи.", get_name());
    return CallbackReturn::SUCCESS;
  }
};

int main(int argc, char * argv[])
{
  rclcpp::init(argc, argv);
  rclcpp::executors::SingleThreadedExecutor executor;

  auto node_1 = std::make_shared<ManagedNode>("managed_node_1");
  auto node_2 = std::make_shared<ManagedNode>("managed_node_2");

  executor.add_node(node_1->get_node_base_interface());
  executor.add_node(node_2->get_node_base_interface());

  executor.spin();
  rclcpp::shutdown();
  return 0;
}""")

add_code(doc, "Лістинг 1.6 – Вихідний код src/lifecycle_supervisor.cpp:", """#include <chrono>
#include <memory>
#include <string>
#include <vector>

#include "rclcpp/rclcpp.hpp"
#include "lifecycle_msgs/msg/transition.hpp"
#include "lifecycle_msgs/srv/change_state.hpp"

using namespace std::chrono_literals;

class LifecycleSupervisor : public rclcpp::Node
{
public:
  LifecycleSupervisor() : Node("lifecycle_supervisor") {}

  bool change_state(const std::string & node_name, std::uint8_t transition_id)
  {
    std::string service_name = "/" + node_name + "/change_state";
    auto client = this->create_client<lifecycle_msgs::srv::ChangeState>(service_name);

    if (!client->wait_for_service(5s)) {
      RCLCPP_ERROR(get_logger(), "Сервіс %s недоступний!", service_name.c_str());
      return false;
    }

    auto request = std::make_shared<lifecycle_msgs::srv::ChangeState::Request>();
    request->transition.id = transition_id;

    auto future = client->async_send_request(request);
    if (rclcpp::spin_until_future_complete(this->get_node_base_interface(), future) ==
        rclcpp::FutureReturnCode::SUCCESS)
    {
      return future.get()->success;
    }
    return false;
  }
};

int main(int argc, char * argv[])
{
  rclcpp::init(argc, argv);
  auto supervisor = std::make_shared<LifecycleSupervisor>();

  std::vector<std::string> nodes = {"managed_node_1", "managed_node_2"};

  for (const auto & name : nodes) {
    RCLCPP_INFO(supervisor->get_logger(), "Конфігурація %s...", name.c_str());
    supervisor->change_state(name, lifecycle_msgs::msg::Transition::TRANSITION_CONFIGURE);
  }

  for (const auto & name : nodes) {
    RCLCPP_INFO(supervisor->get_logger(), "Активація %s...", name.c_str());
    supervisor->change_state(name, lifecycle_msgs::msg::Transition::TRANSITION_ACTIVATE);
  }

  rclcpp::shutdown();
  return 0;
}""")

add_p(doc, "Після запуску managed_nodes_exec та lifecycle_supervisor обидва вузли пройшли конфігурацію та активацію. Перевірка командами ros2 lifecycle get /managed_node_1 та ros2 lifecycle get /managed_node_2 підтвердила стан active [3].")

add_img(doc, "Знімок екрана з 2026-10-15 16-02-11.png", "Рисунок 1.3 – Логи переходів станів managed_node_1 та managed_node_2")
add_img(doc, "Знімок екрана з 2026-10-15 16-02-22.png", "Рисунок 1.4 – Робота lifecycle_supervisor та перевірка активного стану вузлів (active [3])")

# Розділ 3
add_h(doc, "3. Відповіді на контрольні запитання")
add_p(doc, "1. Які переваги та потенційні ризики виникають при запуску кількох вузлів-компонентів у межах одного процесу-виконавця (Executor)?", bold=True)
add_p(doc, "Запуск кількох компонентів в одному процесі суттєво знижує затримки комунікації за рахунок використання спільної пам'яті, але створює ризик: критична помилка в одному компоненті призведе до аварійного завершення всього контейнера та всіх суміжних вузлів.")

add_p(doc, "2. Які обов'язкові вимоги висуваються до C++ класу для можливості його експорту та динамічного завантаження як компонента ROS 2?", bold=True)
add_p(doc, "Клас має успадковувати rclcpp::Node, приймати rclcpp::NodeOptions у конструкторі, бути зареєстрованим макросом RCLCPP_COMPONENTS_REGISTER_NODE у вихідному файлі та компілюватися у спільну бібліотеку (SHARED) із реєстрацією через rclcpp_components_register_nodes у CMakeLists.txt.")

add_p(doc, "3. Чому виклик rclcpp::spin_once() у конструкторі Lifecycle-вузла є небезпечним і до яких критичних помилок це може призвести?", bold=True)
add_p(doc, "У конструкторі об'єкт ще не повністю ініціалізований і не керується зовнішнім std::shared_ptr, тому виклик spin_once() призведе до винятку std::bad_weak_ptr та порушить детерміновану логіку станів LifecycleNode до переходу у стан Unconfigured.")

add_p(doc, "4. Опишіть основні стани та переходи скінченного автомата життєвого циклу керованого вузла (Lifecycle Node).", bold=True)
add_p(doc, "Основні стани: Unconfigured (початковий стан після створення), Inactive (вузол сконфігуровано, але він не виконує активної обробки), Active (основний робочий режим), Finalized (кінцевий стан перед знищенням). Основні переходи між ними: configure (Unconfigured -> Inactive), activate (Inactive -> Active), deactivate (Active -> Inactive), cleanup (Inactive -> Unconfigured) та shutdown (перехід у Finalized).")

# Розділ 4
add_h(doc, "4. Висновки")
add_p(doc, "У ході виконання лабораторної роботи №1 опановано контейнеризацію середовища розробки ROS 2 за допомогою Docker та автоматичного скрипта entrypoint.sh. За допомогою утиліти colcon зібрано пакет lab1_pkg, розроблено та динамічно завантажено вузол-компонент lab1_ns::MyComponent у процес component_container, а також реалізовано два керовані вузли (LifecycleNode) та супервізор, який перевів їх зі стану Unconfigured у стан Active (active [3]).")

out = "Звіт_ЛР1_Турчик_КВ-34.docx"
doc.save(out)
print(f"\nГотово! Файл збережено: {os.path.abspath(out)}")
